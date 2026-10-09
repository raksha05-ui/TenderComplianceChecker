"""Per-page text extraction for PDF and DOCX. Page numbers are list index + 1."""
import io
from dataclasses import dataclass, field

from modules.validators import ValidationError, validate_upload

MIN_CHARS_PER_PAGE = 25  # below this average the file is treated as image-only


@dataclass
class ExtractedDoc:
    name: str
    pages: list            # list[str]; pages[0] is page 1
    needs_ocr: bool = False
    warnings: list = field(default_factory=list)

    @property
    def char_count(self):
        return sum(len(p) for p in self.pages)


def _pdf_pages(name, data):
    import pymupdf
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
        if doc.needs_pass:
            raise ValidationError(f"{name}: the PDF is password protected.")
        return [pg.get_text("text") for pg in doc]
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"{name}: could not read this PDF (corrupted or invalid): {exc}") from exc


def _docx_pages(name, data):
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError(f"{name}: could not read this DOCX (corrupted or invalid): {exc}") from exc
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return ["\n".join(parts)]  # DOCX has no fixed pages: reported as page 1


def extract_document(name, data):
    ext = validate_upload(name, data)
    pages = _pdf_pages(name, data) if ext == ".pdf" else _docx_pages(name, data)
    if not pages:
        raise ValidationError(f"{name}: the document has no pages.")
    doc = ExtractedDoc(name=name, pages=pages)
    if doc.char_count / len(pages) < MIN_CHARS_PER_PAGE:
        doc.needs_ocr = True
        doc.warnings.append(f"{name}: no extractable text - it looks scanned/image-only. "
                            "Run OCR first (e.g. `ocrmypdf in.pdf out.pdf`) and upload the result.")
    if ext == ".docx":
        doc.warnings.append(f"{name}: DOCX has no page numbers; all evidence is reported as page 1.")
    return doc
